"""The downloadable workbook.

Sheets: Summary, each ledger as exported (matched rows shaded, a Status column
added), Matched, To review, and one Unmatched sheet per company. Amounts are
written as numbers, dates as dates, so the sheets sort and sum in Excel.
"""

from __future__ import annotations

import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from .engine import CONFIRMED, REVIEW, Result, rand
from .parse import Ledger, Txn, parse_cents

MONEY_FMT = "#,##0.00"
DATE_FMT = "DD/MM/YYYY"
HEADER_FILL = PatternFill("solid", start_color="F2F2F2", end_color="F2F2F2")
HEADER_FONT = Font(bold=True)
FILL = {
    CONFIRMED: PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE"),
    REVIEW: PatternFill("solid", start_color="FFEB9C", end_color="FFEB9C"),
}
STATUS_LABEL = {CONFIRMED: "Matched", REVIEW: "To review"}
MONEY_HEADERS = {"debit", "credit", "balance", "amount"}


def _sheet_title(name: str, taken: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "", name).strip() or "Sheet"
    title, n = base[:31], 2
    while title.lower() in taken:
        suffix = f" ({n})"
        title, n = base[:31 - len(suffix)] + suffix, n + 1
    taken.add(title.lower())
    return title


def _header(ws: Worksheet, headers: list[str]) -> None:
    for c, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def _money(cell, cents: int | None) -> None:
    cell.value = None if cents is None else cents / 100
    cell.number_format = MONEY_FMT


def _date(cell, when) -> None:
    cell.value = when
    cell.number_format = DATE_FMT


def _autofit(ws: Worksheet) -> None:
    for col in ws.columns:
        longest = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(max(longest + 3, 10), 60)


def _ledger_sheet(ws: Worksheet, ledger: Ledger, status: dict[int, str]) -> None:
    _header(ws, ledger.headers + ["Status"])
    money_cols = {i for i, h in enumerate(ledger.headers) if h.strip().lower() in MONEY_HEADERS}
    for r, row in enumerate(ledger.rows, start=2):
        tier = status.get(r - 2)
        for c, value in enumerate(row[:len(ledger.headers)], start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if c - 1 in money_cols and value.strip():
                try:
                    _money(cell, parse_cents(value))
                except ValueError:
                    pass
            if tier:
                cell.fill = FILL[tier]
        status_cell = ws.cell(row=r, column=len(ledger.headers) + 1, value=STATUS_LABEL.get(tier, ""))
        if tier:
            status_cell.fill = FILL[tier]


def _matches_sheet(ws: Worksheet, matches, name_a: str, name_b: str) -> None:
    _header(ws, ["Amount", f"{name_a} date", f"{name_a} description", f"{name_a} debit",
                 f"{name_a} credit", f"{name_b} date", f"{name_b} description", f"{name_b} debit",
                 f"{name_b} credit", "Days apart"])
    for r, m in enumerate(matches, start=2):
        _money(ws.cell(r, 1), abs(m.a.cents))
        _date(ws.cell(r, 2), m.a.date)
        ws.cell(r, 3, m.a.description)
        _money(ws.cell(r, 4), m.a.debit or None)
        _money(ws.cell(r, 5), m.a.credit or None)
        _date(ws.cell(r, 6), m.b.date)
        ws.cell(r, 7, m.b.description)
        _money(ws.cell(r, 8), m.b.debit or None)
        _money(ws.cell(r, 9), m.b.credit or None)
        ws.cell(r, 10, m.days)


def action_text(t: Txn, other: str) -> str:
    side = "credit" if t.cents > 0 else "debit"
    return f"{other} needs a {side} of {rand(abs(t.cents))} around {t.date:%d/%m/%Y}"


def _unmatched_sheet(ws: Worksheet, txns: list[Txn], other: str) -> None:
    _header(ws, ["Date", "Description", "Debit", "Credit", "Action"])
    for r, t in enumerate(txns, start=2):
        _date(ws.cell(r, 1), t.date)
        ws.cell(r, 2, t.description)
        _money(ws.cell(r, 3), t.debit or None)
        _money(ws.cell(r, 4), t.credit or None)
        ws.cell(r, 5, action_text(t, other))


def _summary_sheet(ws: Worksheet, res: Result) -> None:
    _header(ws, ["", "Debits", "Credits", "Net"])
    for r, ledger in enumerate((res.a, res.b), start=2):
        ws.cell(r, 1, ledger.name)
        _money(ws.cell(r, 2), ledger.debits)
        _money(ws.cell(r, 3), ledger.credits)
        _money(ws.cell(r, 4), ledger.net)
    ws.cell(4, 1, "Difference").font = HEADER_FONT
    _money(ws.cell(4, 4), res.difference)
    ws.cell(4, 4).font = HEADER_FONT
    rows = [
        ("Matched", len(res.confirmed)),
        (f"To review (within {res.tolerance} days)", len(res.review)),
        (f"Unmatched in {res.a.name}", len(res.unmatched_a)),
        (f"Unmatched in {res.b.name}", len(res.unmatched_b)),
    ]
    for r, (label, count) in enumerate(rows, start=6):
        ws.cell(r, 1, label)
        ws.cell(r, 2, count)


def workbook_bytes(res: Result) -> bytes:
    wb = Workbook()
    taken: set[str] = set()
    summary = wb.active
    summary.title = _sheet_title("Summary", taken)
    _summary_sheet(summary, res)

    _ledger_sheet(wb.create_sheet(_sheet_title(res.a.name, taken)), res.a, res.status_a())
    _ledger_sheet(wb.create_sheet(_sheet_title(res.b.name, taken)), res.b, res.status_b())
    _matches_sheet(wb.create_sheet(_sheet_title("Matched", taken)), res.confirmed, res.a.name, res.b.name)
    _matches_sheet(wb.create_sheet(_sheet_title("To review", taken)), res.review, res.a.name, res.b.name)
    _unmatched_sheet(wb.create_sheet(_sheet_title(f"Unmatched {res.a.name}", taken)),
                     res.unmatched_a, res.b.name)
    _unmatched_sheet(wb.create_sheet(_sheet_title(f"Unmatched {res.b.name}", taken)),
                     res.unmatched_b, res.a.name)

    for ws in wb.worksheets:
        _autofit(ws)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def file_name(res: Result) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "_", f"{res.a.name} {res.b.name}").strip("_")
    return f"Loan_recon_{slug or 'export'}.xlsx"
