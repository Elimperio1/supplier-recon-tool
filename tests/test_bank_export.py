import io
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from bankrecon.engine import reconcile
from bankrecon.export import workbook_bytes
from bankrecon.parse import SIDE_CSV, SIDE_SAGE, parse_file

FIX = Path(__file__).parent / "fixtures"


def _result():
    bank = parse_file((FIX / "bank_recon_bank.csv").read_bytes(), SIDE_CSV)
    sage = parse_file((FIX / "bank_recon_sage.csv").read_bytes(), SIDE_SAGE)
    return reconcile(bank, sage)


def _rows(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_workbook_round_trip():
    result = _result()
    wb = load_workbook(io.BytesIO(workbook_bytes(result)))
    assert wb.sheetnames == ["Summary", "Missing in Sage", "Extra in Sage", "Review",
                             "Matched", "CSV marked", "Sage marked"]

    summary = {r[0]: r[1] for r in _rows(wb["Summary"])[1:]}
    assert summary["Difference (CSV minus Sage)"] == 1125.0
    assert summary["Missing in Sage (lines)"] == 3
    assert summary["Extra in Sage (lines)"] == 2
    assert summary["Proof: missing minus extra equals difference"] == "PASS"
    assert isinstance(summary["Compared from"], datetime)

    missing = _rows(wb["Missing in Sage"])
    assert missing[0] == ["Date", "Description", "Amount", "Action"]
    assert [r[1] for r in missing[1:]] == ["0108 honouring fee", "SALES DEPOSIT credit transfer",
                                           "CUSTOMER ONE payment"]
    assert [r[2] for r in missing[1:]] == [-155.0, 2500.0, 1200.0]
    assert all(isinstance(r[0], datetime) for r in missing[1:])

    extra = _rows(wb["Extra in Sage"])
    assert [r[4] for r in extra[1:]] == ["SALES DEPOSIT credit transfer", "STATIONERY debit card purchase"]
    assert [r[6] for r in extra[1:]] == [2500.0, -80.0]

    review = _rows(wb["Review"])
    assert len(review) - 1 == 1 + 2 + 2          # amount only, split of 2, reverse split of 2
    assert {r[1] for r in review[1:]} == {"Amount only", "Split"}


def test_marked_sheets_carry_every_input_row():
    result = _result()
    wb = load_workbook(io.BytesIO(workbook_bytes(result)))
    csv_marked = _rows(wb["CSV marked"])
    sage_marked = _rows(wb["Sage marked"])
    assert len(csv_marked) - 1 == len(result.csv_in) + len(result.csv_out) == 12
    assert len(sage_marked) - 1 == len(result.sage_in) + len(result.sage_out) == 11
    csv_status = [r[-1] for r in csv_marked[1:]]
    assert csv_status[0] == "Outside window"
    assert csv_status.count("Missing in Sage") == 3
    sage_status = [r[-1] for r in sage_marked[1:]]
    assert sage_status.count("Extra in Sage") == 2
    assert sage_status[-1] == "Outside window"
    # Rows keep their file order and line numbers.
    assert [r[0] for r in csv_marked[1:]] == list(range(6, 18))
