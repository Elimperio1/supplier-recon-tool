"""Loan recon engine tests. Every figure and name here is invented."""

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from loanrecon.engine import CONFIRMED, REVIEW, rand, reconcile
from loanrecon.export import workbook_bytes
from loanrecon.parse import ParseError, parse_cents, parse_ledger

SAGE_A = """﻿Alpha Holdings (Pty) Ltd
Loan Account: Beta Trading,,,,
Account Date,Reference,Description,Debit,Credit,Balance
01/03/2026,LN1,Loan advanced,"10,000.00",,"10,000.00"
05/03/2026,LN2,Second advance,2500.00,,"12,500.00"
10/03/2026,LN3,Repayment received,,1000.00,"11,500.00"
15/03/2026,LN4,Sign-swapped repayment,-300.00,,"11,200.00"
20/03/2026,LN5,Unmatched advance,750.00,,"11,950.00"
,,Total,"13,250.00","1,300.00",
"""

PASTEL_B = """Date,Description,Debit,Credit
01/03/2026,Loan from Alpha,,10000.00
07/03/2026,Loan from Alpha late,,2500.00
10/03/2026,Repaid Alpha,1000.00,
15/03/2026,Repaid Alpha small,300.00,
25/03/2026,Interest not in Alpha,,120.00
"""


def ledgers(tolerance=2):
    a = parse_ledger(SAGE_A.encode("utf-8"), "Sage", "Alpha")
    b = parse_ledger(PASTEL_B.encode("utf-8"), "Pastel", "Beta")
    return a, b, reconcile(a, b, tolerance)


@pytest.mark.parametrize("text,cents", [
    ("1,234.50", 123450), ("(500.00)", -50000), ("", 0), ("R 1 000", 100000),
    ("-12.3", -1230), ("12.30-", -1230), (" 1 000.05", 100005),
])
def test_parse_cents(text, cents):
    assert parse_cents(text) == cents


def test_parse_cents_rejects_garbage():
    with pytest.raises(ValueError):
        parse_cents("12.3.4")


def test_parse_finds_header_and_skips_footer():
    a, b, _ = ledgers()
    assert len(a.txns) == 5 and len(b.txns) == 5
    # Total row has no date and is not a transaction.
    assert all(t.description != "Total" for t in a.txns)


def test_negative_debit_is_a_credit():
    a, _, _ = ledgers()
    swapped = next(t for t in a.txns if t.description.startswith("Sign-swapped"))
    assert swapped.cents == -30000 and swapped.credit == 30000


def test_tiers_and_unmatched():
    _, _, res = ledgers(tolerance=2)
    assert {m.a.description: m.tier for m in res.matches} == {
        "Loan advanced": CONFIRMED,
        "Second advance": REVIEW,             # two days apart
        "Repayment received": CONFIRMED,
        "Sign-swapped repayment": CONFIRMED,
    }
    assert [t.description for t in res.unmatched_a] == ["Unmatched advance"]
    assert [t.description for t in res.unmatched_b] == ["Interest not in Alpha"]


def test_zero_tolerance_leaves_late_entry_unmatched():
    _, _, res = ledgers(tolerance=0)
    assert not res.review
    assert len(res.unmatched_a) == 2


def test_difference_equals_unmatched_sum():
    a, b, res = ledgers()
    assert a.net == 1000000 + 250000 - 100000 - 30000 + 75000
    assert b.net == -1000000 - 250000 + 100000 + 30000 - 12000
    assert res.difference == sum(t.cents for t in res.unmatched_a + res.unmatched_b) == 63000


def test_nearest_date_wins_in_review_pass():
    a = parse_ledger(b"Date,Description,Debit,Credit\n10/03/2026,X,100,\n", "Pastel", "A")
    b = parse_ledger(b"Date,Description,Debit,Credit\n08/03/2026,far,,100\n11/03/2026,near,,100\n",
                     "Pastel", "B")
    res = reconcile(a, b, 3)
    assert res.matches[0].b.description == "near"


def test_wrong_format_hints_the_right_one():
    with pytest.raises(ParseError, match="Pastel export"):
        parse_ledger(PASTEL_B.encode(), "Sage", "Beta")


def test_missing_column_is_a_clean_error():
    with pytest.raises(ParseError, match="'Credit'"):
        parse_ledger(b"Date,Description,Debit\n01/03/2026,x,1\n", "Pastel", "B")


def test_unreadable_amount_is_reported():
    led = parse_ledger(b"Date,Description,Debit,Credit\n01/03/2026,ok,5,\n02/03/2026,bad,abc,\n",
                       "Pastel", "B")
    assert len(led.txns) == 1 and "unreadable Debit or Credit" in led.warnings[0]


def test_rand_house_format():
    assert rand(123456789) == "R 1 234 567.89"
    assert rand(-5000) == "-R 50.00"


def test_workbook_numbers_dates_and_sheet_names():
    a, b, res = ledgers()
    a.name = b.name = "Same/Name"          # duplicate, with a character Excel forbids
    wb = load_workbook(BytesIO(workbook_bytes(res)))
    assert wb.sheetnames[:3] == ["Summary", "SameName", "SameName (2)"]
    assert wb["Summary"]["D4"].value == pytest.approx(630.00)
    matched = wb["Matched"]
    assert matched["A2"].value == pytest.approx(10000.0)
    assert matched["B2"].value.date() == date(2026, 3, 1)
    ledger = wb["SameName"]
    assert ledger["D2"].value == pytest.approx(10000.0)       # Debit as a number
    assert ledger.cell(2, ledger.max_column).value == "Matched"
