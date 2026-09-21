"""Excel workbook of a reconciliation result (openpyxl, no pandas).

Sheets: Summary, Missing in Sage, Extra in Sage, Review, Matched, CSV marked,
Sage marked. The two marked sheets carry every input row with a Status column,
so the accountant can work from the full list. Amounts are written as numbers
with two decimals and dates as real dates.
"""

from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .engine import (STATUS_AMOUNT, STATUS_DATE, STATUS_EXTRA, STATUS_MATCHED,
                     STATUS_MISSING, STATUS_OUTSIDE, STATUS_SPLIT, Match, Result, Txn)

FILL_RED = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
FILL_AMBER = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
FILL_GREEN = PatternFill(start_color="E7F5EC", end_color="E7F5EC", fill_type="solid")
FILL_GREY = PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid")
FILL_HEAD = PatternFill(start_color="1D1D1F", end_color="1D1D1F", fill_type="solid")

STATUS_FILL = {
    STATUS_MATCHED: FILL_GREEN,
    STATUS_AMOUNT: FILL_AMBER,
    STATUS_SPLIT: FILL_AMBER,
    STATUS_DATE: FILL_AMBER,
    STATUS_MISSING: FILL_RED,
    STATUS_EXTRA: FILL_RED,
    STATUS_OUTSIDE: FILL_GREY,
}

AMOUNT_FORMAT = "#,##0.00;-#,##0.00"
DATE_FORMAT = "dd/mm/yyyy"


def _money(cents: int) -> float:
    # Display only. The engine never compares floats; Excel wants a number.
    return cents / 100


def _write_table(ws, headers: list[str], rows: list[list], fills=None) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = FILL_HEAD
        cell.alignment = Alignment(vertical="center")
    for i, row in enumerate(rows):
        ws.append(row)
        if fills and fills[i] is not None:
            for cell in ws[ws.max_row]:
                cell.fill = fills[i]
    for col_idx, header in enumerate(headers, start=1):
        letter = get_column_letter(col_idx)
        for cell in ws[letter][1:]:
            if isinstance(cell.value, date):
                cell.number_format = DATE_FORMAT
            elif isinstance(cell.value, float):
                cell.number_format = AMOUNT_FORMAT
    _autofit(ws)
    ws.freeze_panes = "A2"


def _autofit(ws) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            text = f"{cell.value:%d/%m/%Y}" if isinstance(cell.value, date) else str(cell.value)
            widths[cell.column] = max(widths.get(cell.column, 0), len(text))
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 10), 60)


def _match_rows(matches: list[Match]) -> list[list]:
    rows: list[list] = []
    for n, m in enumerate(matches, start=1):
        for c in m.csv_txns:
            for s in m.sage_txns:
                rows.append([n, m.tier, c.date, c.description, _money(c.cents),
                             s.date, s.reference, s.description, s.comment, _money(s.cents), m.note])
    return rows


MATCH_HEADERS = ["Match", "Tier", "CSV date", "CSV description", "CSV amount",
                 "Sage date", "Sage reference", "Sage description", "Sage comment",
                 "Sage total", "Note"]


def workbook_bytes(result: Result) -> bytes:
    wb = Workbook()
    status = result.status_map()

    ws = wb.active
    ws.title = "Summary"
    summary = [
        ["Compared from", result.window.start],
        ["Compared to", result.window.end],
        ["Date tolerance (days)", result.tolerance_days],
        ["CSV lines in window", len(result.csv_in)],
        ["Sage lines in window", len(result.sage_in)],
        ["CSV total", _money(result.csv_total)],
        ["Sage total", _money(result.sage_total)],
        ["Difference (CSV minus Sage)", _money(result.difference)],
        ["Matched CSV lines", result.matched_csv_lines],
        ["Matched Sage lines", result.matched_sage_lines],
        ["Missing in Sage (lines)", len(result.missing_in_sage)],
        ["Missing in Sage (total)", _money(result.missing_total)],
        ["Extra in Sage (lines)", len(result.extra_in_sage)],
        ["Extra in Sage (total)", _money(result.extra_total)],
        ["Matches to review", len(result.review)],
        ["CSV lines outside window", len(result.csv_out)],
        ["Sage lines outside window", len(result.sage_out)],
        ["CSV rows skipped", len(result.csv_skipped)],
        ["Sage rows skipped", len(result.sage_skipped)],
        ["Proof: missing minus extra equals difference", "PASS" if result.proof_ok else "FAIL"],
    ]
    _write_table(ws, ["Item", "Value"], summary)

    ws = wb.create_sheet("Missing in Sage")
    _write_table(ws, ["Date", "Description", "Amount", "Action"],
                 [[t.date, t.description, _money(t.cents), "Capture in Sage"]
                  for t in result.missing_in_sage],
                 fills=[FILL_RED] * len(result.missing_in_sage))

    ws = wb.create_sheet("Extra in Sage")
    _write_table(ws, ["Date", "Type", "Selection", "Reference", "Description", "Comment",
                      "Total", "Reconciled", "Action"],
                 [[t.date, t.txn_type, t.selection, t.reference, t.description, t.comment,
                   _money(t.cents), t.reconciled, "Not on the bank CSV, check in Sage"]
                  for t in result.extra_in_sage],
                 fills=[FILL_RED] * len(result.extra_in_sage))

    ws = wb.create_sheet("Review")
    review_rows = _match_rows(result.review)
    _write_table(ws, MATCH_HEADERS, review_rows, fills=[FILL_AMBER] * len(review_rows))

    ws = wb.create_sheet("Matched")
    _write_table(ws, MATCH_HEADERS, _match_rows(result.matches))

    ws = wb.create_sheet("CSV marked")
    csv_all: list[Txn] = sorted(result.csv_in + result.csv_out, key=lambda t: t.row)
    _write_table(ws, ["Row", "Date", "Description", "Amount", "Status"],
                 [[t.row, t.date, t.description, _money(t.cents), status[t.key]] for t in csv_all],
                 fills=[STATUS_FILL[status[t.key]] for t in csv_all])

    ws = wb.create_sheet("Sage marked")
    sage_all: list[Txn] = sorted(result.sage_in + result.sage_out, key=lambda t: t.row)
    _write_table(ws, ["Row", "Date", "Type", "Selection", "Reference", "Description", "Comment",
                      "Total", "Reconciled", "Status"],
                 [[t.row, t.date, t.txn_type, t.selection, t.reference, t.description, t.comment,
                   _money(t.cents), t.reconciled, status[t.key]] for t in sage_all],
                 fills=[STATUS_FILL[status[t.key]] for t in sage_all])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
